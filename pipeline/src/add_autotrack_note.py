import openpyxl
from openpyxl.styles import Font, Alignment
from datetime import date

wb = openpyxl.load_workbook('cricket_tracking.xlsx')
ws2 = wb['Notes']
BASE_FONT = Font(name='Arial')
new_notes = [
    (str(date.today()), "First real attempt at getting the computer to find the ball on its own, instead of us "
                          "reading pixels by hand. Simple idea: look at what moved between one frame and the "
                          "next AND the one after that (so it has to be moving consistently, not a one-off flicker), "
                          "then out of everything that's moving, pick whatever is closest to where a straight line "
                          "would predict the ball to be next. Tested on C4 against our own careful hand-readings."),
    (str(date.today()), "RESULT: it tracked the ball on its own, matching our hand-tracking to within about 7-8cm "
                          "the whole way from release almost to the bounce - genuinely promising for something "
                          "built in a few minutes with no training data. But it got confused at the exact moment "
                          "the ball bounces (picked frame 105 instead of the real bounce at frame 104), because "
                          "that's the one moment the ball suddenly changes direction, and our 'predict a straight "
                          "line' logic breaks exactly then. At the correct bounce frame itself, the auto-tracked "
                          "position was still only about 4cm off the true sideways answer - almost as good as our "
                          "hand-tracked number. Next step: build a proper bounce-detector that specifically looks "
                          "for the moment downward motion flips to upward, rather than trusting the general "
                          "tracker's guess at where it lands."),
]
for row in new_notes:
    ws2.append(row)
for row_i in range(ws2.max_row - len(new_notes) + 1, ws2.max_row + 1):
    ws2.cell(row=row_i, column=1).font = BASE_FONT
    ws2.cell(row=row_i, column=2).font = BASE_FONT
    ws2.cell(row=row_i, column=2).alignment = Alignment(wrap_text=True)
wb.save('cricket_tracking.xlsx')
print('done')

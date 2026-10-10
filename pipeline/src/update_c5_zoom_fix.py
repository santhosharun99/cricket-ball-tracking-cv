import openpyxl
from openpyxl.styles import Font, Alignment
from datetime import date

wb = openpyxl.load_workbook('cricket_tracking.xlsx')
ws = wb['Our_Predictions']
# C5 row = 6
ws['F6'] = -0.308   # pred_bounce_y (zoom-corrected, frame84)
ws['K6'] = -0.276   # pred_stump_y (zoom-corrected, frame85, right after bounce)
# clear the old bad stump_z guess since it was built on the flat-scale method - leave blank rather than wrong
ws['L6'] = None      # pred_stump_z - retract old impossible value, not yet re-derived properly
wb.save('cricket_tracking.xlsx')

ws2 = wb['Notes']
BASE_FONT = Font(name='Arial')
new_notes = [
    (str(date.today()), "Checked C5 for the same camera-zoom effect found in C4 - yes, same thing happens, just "
                          "slower (about 0.94 pixels of stump height per frame here, versus ~5.6 in C4 - each clip's "
                          "camera operator zooms at a different pace). Re-measured C5's stumps at two points during "
                          "the run-up (frame 42 and frame 60) to work out this clip's own zoom rate, then applied "
                          "the same per-frame ruler fix used for C4."),
    (str(date.today()), "RESULT for C5 with the fix: bounce point sideways position comes out at -0.308m (real "
                          "answer -0.33m, off by just 2cm - previously this method gave -0.499m, way off). The "
                          "ball's position right after the bounce comes out at -0.276m, versus the real stump_y of "
                          "-0.279m - a 0.3cm difference, essentially exact. Both clips now check out well on the "
                          "sideways (off/leg) numbers once the zoom is accounted for properly. Cleared the old "
                          "stump_z (height) guess for C5 since it was built on the broken flat-ruler method - not "
                          "yet re-derived with the fix, left blank rather than keep a wrong number."),
]
for row in new_notes:
    ws2.append(row)
for row_i in range(ws2.max_row - len(new_notes) + 1, ws2.max_row + 1):
    ws2.cell(row=row_i, column=1).font = BASE_FONT
    ws2.cell(row=row_i, column=2).font = BASE_FONT
    ws2.cell(row=row_i, column=2).alignment = Alignment(wrap_text=True)
wb.save('cricket_tracking.xlsx')
print('saved, notes rows now', ws2.max_row)

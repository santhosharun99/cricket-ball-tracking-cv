import openpyxl
from openpyxl.styles import Font, Alignment
from datetime import date

wb = openpyxl.load_workbook('cricket_tracking.xlsx')
ws2 = wb['Notes']
BASE_FONT = Font(name='Arial')
new_notes = [
    (str(date.today()), "Tried to fix the release-position number properly, as asked. First thing to fix: our "
                          "left-right 'centre line' at the release frame was only ever extrapolated backwards from "
                          "two points measured 11 frames later, near the bounce - a shaky stretch. So this time we "
                          "went back to the actual release frame and directly measured the real stumps sitting "
                          "there (two of the three are visible past the batter's legs, giving us a real pixel "
                          "reading instead of a guess)."),
    (str(date.today()), "RESULT: it did NOT improve things - using the directly-measured stumps gave a WORSE "
                          "sideways release number than our original stretched-line guess, not a better one. That "
                          "tells us something useful even though it's a negative result: the leftover error wasn't "
                          "really coming from a shaky centre line after all. The more likely real cause is what we "
                          "flagged before - at the exact instant of release the ball is up near the bowler's head, "
                          "about 2 metres off the ground, not sitting on the pitch, and a TV camera looking at an "
                          "angle (not straight overhead) will show something raised up in the air slightly "
                          "shifted sideways compared to a ground-level ruler. Properly correcting for that needs "
                          "to know things about the camera itself - exactly how high it's mounted, how far back, "
                          "and how steeply it's tilted down - and none of that can be read off a normal broadcast "
                          "video with any confidence. Rather than bolt on a correction built on guesses about the "
                          "camera, we're leaving the release sideways-position at its current accuracy (about 7-8cm "
                          "off) and calling this one closed for now - it's already usable, and chasing the last few "
                          "centimetres here needs information we don't have access to."),
]
for row in new_notes:
    ws2.append(row)
for row_i in range(ws2.max_row - len(new_notes) + 1, ws2.max_row + 1):
    ws2.cell(row=row_i, column=1).font = BASE_FONT
    ws2.cell(row=row_i, column=2).font = BASE_FONT
    ws2.cell(row=row_i, column=2).alignment = Alignment(wrap_text=True)
wb.save('cricket_tracking.xlsx')
print('done')

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from datetime import date

wb = openpyxl.load_workbook('cricket_tracking.xlsx')
FONT = "Arial"
BASE_FONT = Font(name=FONT)
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")

if 'Score_Estimate' in wb.sheetnames:
    del wb['Score_Estimate']
ws = wb.create_sheet('Score_Estimate', 3)

ws['A1'] = "Rough scorecard as of now - our best honest guess, using the real subnet weights"
ws['A1'].font = Font(name=FONT, bold=True)
ws.merge_cells('A1:F1')

ws.append([])
ws.append(["IMPORTANT CAVEAT:"])
ws.cell(row=ws.max_row, column=1).font = Font(name=FONT, bold=True)
cav = [
    "We don't actually know the exact scoring formula the validator uses (is it strict pass/fail within a "
    "tolerance, or does 'closer counts for more' on a sliding scale?). So the numbers below are OUR estimate of "
    "how much credit each field probably earns given how close we usually land - not a guaranteed number.",
    "Also based on real testing on only 2 of the 5 example clips for the harder fields (bounce_x, kph) - the "
    "other 3 haven't been run through our pipeline yet, so real performance could shift once we do.",
]
for s in cav:
    ws.append([s]); ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(f'A{ws.max_row}:F{ws.max_row}')

ws.append([])
headers = ["field(s)", "weight (%)", "our credit estimate", "points earned", "status", "why"]
ws.append(headers)
hdr_row = ws.max_row

fields = [
    ("bounce_x", 19, "55%", 10.5, "PARTIAL", "real working method now (crease + timing + speed) but only tested on 2 clips - nailed one (4-7cm off), rough on the other (33cm off)"),
    ("kph", 16, "60%", 9.6, "PARTIAL", "great when the speed-gun graphic is shown; own back-up method decent but leans on already having a bounce point"),
    ("stump_y", 14, "90%", 12.6, "GOOD", "our best, most consistent number - 1-4cm off"),
    ("deviation", 10, "0%", 0.0, "NOT BUILT", "haven't built the calc yet, even though we already have both pieces it needs"),
    ("swing_angle", 8, "0%", 0.0, "NOT BUILT", "same - pieces are there, calc not built"),
    ("stump_z", 8, "45%", 3.6, "PARTIAL", "usable but the softest of our working numbers"),
    ("release_y", 2, "80%", 1.6, "GOOD", "7-8cm off"),
    ("release_z", 2, "0%", 0.0, "MISSING", "never calculated ourselves, only ever used the answer key"),
    ("bounce_y", 2, "90%", 1.8, "GOOD", "2-4cm off, very solid"),
    ("impact_x", 2, "0%", 0.0, "MISSING", "not attempted - same root problem as bounce_x"),
    ("impact_y", 2, "80%", 1.6, "GOOD", "1-4cm off"),
    ("impact_z", 2, "45%", 0.9, "PARTIAL", "soft, like stump_z"),
    ("interception_distance", 2, "0%", 0.0, "MISSING", "needs bounce_x and impact_x, both shaky/missing"),
    ("inningsid + overid + ball_in_over", 6, "80%", 4.8, "GOOD", "scorecard-reading (OCR) proved accurate on the one clip we checked"),
    ("ballid", 1, "70%", 0.7, "OK", "should be easy bookkeeping once over/ball are known"),
    ("xlsx_overs + scorecard_overs", 2, "75%", 1.5, "GOOD", "same scorecard-reading approach"),
    ("runs", 1, "85%", 0.9, "GOOD", "scorecard delta matched exactly on our one test"),
    ("wickets", 1, "85%", 0.9, "GOOD", "scorecard delta matched exactly on our one test"),
]
for f in fields:
    ws.append(f)

fill_map = {"GOOD": GREEN, "OK": GREEN, "PARTIAL": YELLOW, "MISSING": RED, "NOT BUILT": RED}
for i, f in enumerate(fields):
    row_i = hdr_row+1+i
    for c in range(1,7):
        ws.cell(row=row_i, column=c).font = BASE_FONT
        ws.cell(row=row_i, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=row_i, column=5).fill = fill_map[f[4]]

for col_idx, h in enumerate(headers, 1):
    c = ws.cell(row=hdr_row, column=col_idx)
    c.font = HEADER_FONT; c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", wrap_text=True)

total_row = ["TOTAL", 100, "", round(sum(f[3] for f in fields),1), "", "~51 out of 100, roughly half marks"]
ws.append(total_row)
for c in range(1,7):
    cell = ws.cell(row=ws.max_row, column=c)
    cell.font = Font(name=FONT, bold=True)
    cell.fill = PatternFill("solid", fgColor="D9D9D9")
    cell.alignment = Alignment(wrap_text=True)

note = ws.cell(row=ws.max_row+2, column=1,
    value="Bottom line: if we plugged our current pipeline in today, our best honest guess is somewhere around "
          "half marks - roughly 50%. The single biggest lever isn't accuracy at all, it's that swing_angle (8%) "
          "and deviation (10%) - 18 points combined - are sitting at zero purely because we haven't wired up the "
          "calculation yet, even though we already have every number both formulas need (release_y, bounce_y, "
          "stump_y/impact_y). That alone could take us from ~51% to ~69% with no new tracking work at all. After "
          "that, tightening bounce_x and kph (35% of the score between them) by testing and fixing on all 5 clips "
          "instead of 2 is the next biggest gain.")
note.font = Font(name=FONT, italic=True, size=9, color="808080")
note.alignment = Alignment(wrap_text=True)
ws.merge_cells(f'A{note.row}:F{note.row}')
ws.row_dimensions[note.row].height = 100

widths = [30, 12, 16, 14, 12, 55]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A5"

wb.save('cricket_tracking.xlsx')

ws2 = wb['Notes']
new_note = (str(date.today()), "Santhosh asked for an honest current scorecard estimate. Used the REAL subnet "
    "weights (from Score_Dashboard.pdf) against our actual accuracy on each field and worked out a rough total: "
    "about 51 out of 100 as things stand today. Added a new Score_Estimate sheet with the full breakdown and the "
    "big caveat that we don't know the validator's exact scoring formula (pass/fail vs sliding scale) or how we'd "
    "do on the 3 example clips we haven't run yet - this is our best honest guess, not a guarantee. Biggest single "
    "opportunity found: swing_angle and deviation (18% combined) are worth zero right now purely because we "
    "haven't coded the formula yet, even though we already have every number they need - that's free points "
    "sitting on the table, not a hard tracking problem.")
ws2.append(new_note)
ws2.cell(row=ws2.max_row, column=1).font = Font(name=FONT)
ws2.cell(row=ws2.max_row, column=2).font = Font(name=FONT)
ws2.cell(row=ws2.max_row, column=2).alignment = Alignment(wrap_text=True)
wb.save('cricket_tracking.xlsx')
print('done')

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from datetime import date

wb = openpyxl.load_workbook('cricket_tracking.xlsx')
FONT = "Arial"
BASE_FONT = Font(name=FONT)
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")

if 'Calibration_Upgrade_Attempt' in wb.sheetnames:
    del wb['Calibration_Upgrade_Attempt']
ws = wb.create_sheet('Calibration_Upgrade_Attempt', 3)

ws['A1'] = "Tried to improve calibration precision (idea borrowed from a real Basque Pelota ball-tracking paper) - it made things worse, not better"
ws['A1'].font = Font(name=FONT, bold=True)
ws.merge_cells('A1:F1')

ws.append([])
ws.append(["Where the idea came from:"])
ws.cell(row=ws.max_row, column=1).font = Font(name=FONT, bold=True)
p0 = [
    "Santhosh shared a real, patented ball-bounce-tracking system used on live Basque Pelota TV broadcasts "
    "(Labayen et al., Vicomtech). Their big advantage over us: they calibrate their camera properly using a "
    "checkerboard pattern laid on the court before the match, giving a precise, well-anchored map from pixels to "
    "real-world position. That's why they get accuracy of a few millimetres even 9 metres from the camera - much "
    "tighter than the wobble we've been fighting.",
    "We can't lay down a checkerboard on footage that's already been filmed, but the lesson still applies: our "
    "calibration is only as good as the reference points we feed it. We were using 3 points (stump base, release, "
    "bounce). We tried adding a 4th - the height of the stumps themselves (base to top, a known 0.711m) - to also "
    "correct for how height above the ground shifts a point sideways in the picture, which our first calibration "
    "completely ignored.",
]
for s in p0:
    ws.append([s]); ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(f'A{ws.max_row}:F{ws.max_row}')

ws.append([])
headers = ["version", "swing_angle result (C4)", "true value", "verdict"]
ws.append(headers)
hdr_row = ws.max_row
rows = [
    ("3-point calibration (release, bounce, stump base) - no height correction", "+0.02°", "-0.16°", "closest so far - kept as our current best"),
    ("4-point calibration (added stump top/height) - attempted upgrade", "+0.29°", "-0.16°", "worse - rejected, not used"),
]
for r in rows:
    ws.append(r)
for i in range(len(rows)):
    row_i = hdr_row+1+i
    for c in range(1,5):
        ws.cell(row=row_i, column=c).font = BASE_FONT
        ws.cell(row=row_i, column=c).alignment = Alignment(wrap_text=True, vertical="top")
for col_idx, h in enumerate(headers, 1):
    c = ws.cell(row=hdr_row, column=col_idx)
    c.font = HEADER_FONT; c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", wrap_text=True)

ws.append([])
ws.append(["Why it backfired:"])
ws.cell(row=ws.max_row, column=1).font = Font(name=FONT, bold=True)
p1 = [
    "The stumps are only about 12 pixels tall in our footage (0.711m of real height squeezed into a tiny span "
    "on screen), compared to well over 100 pixels for the release-to-bounce distances we already use. Measuring "
    "something over such a short pixel span is like trying to time a race with a stopwatch that only ticks once "
    "a second - a single pixel of measurement wobble turns into a huge swing in the answer. Adding this point "
    "didn't sharpen the calibration, it poisoned it with a noisy new number.",
    "We also went looking for the crease lines painted on the pitch (a wider, more reliable reference, like the "
    "Pelota paper effectively uses) as a better alternative to the stump height. No luck - in our footage the "
    "crease markings are too faint/worn to pin down precisely, and in one search we mostly found a bowler's "
    "discarded cap lying on the pitch instead of a painted line.",
]
for s in p1:
    ws.append([s]); ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(f'A{ws.max_row}:F{ws.max_row}')

note = ws.cell(row=ws.max_row+2, column=1,
    value="Bottom line: kept the original 3-point calibration (it's still our best result) and did NOT adopt this "
          "4-point version. The real takeaway from the Pelota paper stands as a concept - better reference points "
          "would genuinely help - but standard broadcast cricket footage doesn't hand us a good enough extra "
          "landmark to act on that concept right now, at least not with manual pixel-reading. Getting further here "
          "would need either sharper source video or a properly surveyed camera set-up, neither of which we have "
          "for these 5 example clips. Recommend moving on from chasing more swing_angle/deviation precision and "
          "instead spending effort on the 3 untested clips, where the automatic ball-detector can do the frame-by-"
          "frame tracking work for bounce_x and kph (35% of the score, vs 18% for swing_angle+deviation combined) "
          "without more manual pixel-hunting.")
note.font = Font(name=FONT, italic=True, size=9, color="808080")
note.alignment = Alignment(wrap_text=True)
ws.merge_cells(f'A{note.row}:F{note.row}')
ws.row_dimensions[note.row].height = 110

widths = [45, 22, 12, 40]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

wb.save('cricket_tracking.xlsx')

ws2 = wb['Notes']
new_notes = [
    (str(date.today()), "Santhosh said to go ahead with the calibration upgrade idea from the Basque Pelota "
        "paper (their trick: use a proper checkerboard-based calibration instead of a few guessed points). "
        "Tried the closest equivalent we could manage on already-filmed footage: added a 4th calibration point "
        "(stump height, base to top) so the model would also correct for how far above the ground a point is, "
        "which our first calibration ignored entirely."),
    (str(date.today()), "Result: it made swing_angle WORSE, not better (C4 went from +0.02 to +0.29 vs true "
        "-0.16). Root cause - the stumps are only about 12 pixels tall in our footage, so measuring their height "
        "in pixels is extremely noisy, and that noise poisoned the whole calibration once fed in. Also tried to "
        "find the painted crease lines as a sturdier alternative reference point (which is closer to what the "
        "Pelota paper actually does) - the crease markings in our footage are too faint/worn to pinpoint "
        "reliably, so that option is off the table with this footage."),
    (str(date.today()), "Decision: kept the original 3-point calibration (our best result so far) and rejected "
        "the 4-point version. Recommending we stop chasing more swing_angle/deviation precision for now - it's "
        "hit a real wall given what our footage can offer - and instead spend effort testing the trained "
        "ball-detector on the 3 example clips we haven't touched yet, since bounce_x and kph together are worth "
        "almost twice as much (35%) as swing_angle and deviation combined (18%)."),
]
for row in new_notes:
    ws2.append(row)
for row_i in range(ws2.max_row - len(new_notes) + 1, ws2.max_row + 1):
    ws2.cell(row=row_i, column=1).font = BASE_FONT
    ws2.cell(row=row_i, column=2).font = BASE_FONT
    ws2.cell(row=row_i, column=2).alignment = Alignment(wrap_text=True)
wb.save('cricket_tracking.xlsx')
print('done')

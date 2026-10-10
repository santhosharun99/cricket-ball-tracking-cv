import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from datetime import date

wb = openpyxl.load_workbook('cricket_tracking.xlsx')
FONT = "Arial"
BASE_FONT = Font(name=FONT)
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")

ws = wb['Our_Predictions']
# add release_y columns if not present - check headers
headers = [c.value for c in ws[1]]
if 'pred_release_y' not in headers:
    col = len(headers) + 1
    ws.cell(row=1, column=col, value='pred_release_y')
    ws.cell(row=1, column=col).font = HEADER_FONT
    ws.cell(row=1, column=col).fill = HEADER_FILL
    # C4 is row5, C5 is row6 (established earlier in this sheet)
    ws.cell(row=5, column=col, value=-0.647)
    ws.cell(row=6, column=col, value=-0.665)
wb.save('cricket_tracking.xlsx')

ws2 = wb['Notes']
new_notes = [
    (str(date.today()), "Santhosh's question: have we worked out where the bowler actually let go of the ball, "
                          "using the same stump-ruler trick as everywhere else? Answer: not until now - we'd only "
                          "ever used the OFFICIAL answer for the ball's release height, and just guessed a typical "
                          "distance for how far in front of the crease the bowler releases. So we went and "
                          "calculated the sideways release position ourselves the same way we do everything else."),
    (str(date.today()), "RESULT for sideways release position: C4 comes out at -0.647m vs the true -0.57m (7.7cm "
                          "off); C5 comes out at -0.665m vs the true -0.596m (6.9cm off). Both correct in direction "
                          "and roughly the right size, but noticeably less accurate than our bounce/near-bat numbers "
                          "(which were within 1-4cm). Reason: our sideways 'centre line' (where the stumps sit "
                          "left-right) was only pinned down using two measurements taken right near the bounce - "
                          "stretching that same line all the way back to the release frame (10-11 frames earlier) "
                          "means any tiny wobble in that line gets magnified over a longer stretch. Fixable by "
                          "taking a couple of real stump-position readings earlier in the run-up too, instead of "
                          "stretching a 2-point line so far."),
    (str(date.today()), "Could NOT do the same for release HEIGHT the way we did for the bounce, and it's worth "
                          "recording why: our height trick needs two known 'ball is touching the ground here' spots "
                          "to work from (we used the stumps and the bounce point) - but the bowler releases the "
                          "ball from all the way at the OTHER end of the pitch, nowhere near either of those two "
                          "spots, so stretching our existing ground-level line out that far isn't trustworthy. A "
                          "possible fix: use the bowler's own front foot touching the ground at the moment of "
                          "release as a THIRD known ground spot near that end of the pitch - not attempted yet, and "
                          "likely to be rough since the bowler is small and far away in frame at that point, but "
                          "worth trying next if we want a self-calculated release height instead of using the "
                          "answer key's number. Also still don't have a way to work out how far in front of the "
                          "crease the bowler released it (release depth) - currently just assuming a typical "
                          "distance, not measuring it."),
]
for row in new_notes:
    ws2.append(row)
for row_i in range(ws2.max_row - len(new_notes) + 1, ws2.max_row + 1):
    ws2.cell(row=row_i, column=1).font = BASE_FONT
    ws2.cell(row=row_i, column=2).font = BASE_FONT
    ws2.cell(row=row_i, column=2).alignment = Alignment(wrap_text=True)
wb.save('cricket_tracking.xlsx')
print('done')

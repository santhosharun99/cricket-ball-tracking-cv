import openpyxl
from openpyxl.styles import Font, Alignment
from datetime import date

wb = openpyxl.load_workbook('cricket_tracking.xlsx')
ws2 = wb['Notes']
BASE_FONT = Font(name='Arial')
new_notes = [
    (str(date.today()), "Santhosh shared a 2025 IEEE Access paper (Abbas et al., 'Front Pitch View Shot Extraction "
                          "and Ball Tracking in Cricket Using Deep-Learning') that tackles almost this exact "
                          "problem. Good confirmation of our approach: to work out how far the ball pitches, they "
                          "do the SAME trick we do - use a known real-world size as a ruler and scale it by how "
                          "zoomed-in the camera picture is (they use the batter's own height as the ruler; we used "
                          "the fixed stump height, which is more reliable since it never changes person to person). "
                          "They say outright that a normal broadcast camera has no real depth information and "
                          "everyone has to approximate this way - so the workaround we built isn't a shortcut, it's "
                          "the standard approach even in published research."),
    (str(date.today()), "They also show this task is genuinely hard even with a full deep-learning pipeline: their "
                          "own length classifier (full/good/short) only gets it right about 70% of the time on "
                          "their test set. So our classification isn't behind for a hand-built prototype - it's "
                          "just that this problem doesn't have an easy, perfect answer for anyone. One thing to "
                          "flag: their exact length boundaries differ from the ones we used (they use under 6m = "
                          "full, 6-8.5m = good, over 8.5m = short, vs our 2/4/8m bands) - there's no single agreed "
                          "definition across sources, which is worth knowing if we ever get asked to justify the "
                          "exact cutoffs."),
    (str(date.today()), "Two useful takeaways for next steps: (1) they publish an annotated ball-tracking dataset "
                          "and a trained ball-detector model on GitHub (github.com/KumailAbb/FPVShotDetection) - "
                          "worth checking whether it could take over the manual frame-by-frame pixel reading we've "
                          "been doing by hand, which has been our biggest time sink. (2) for lining up the pitch's "
                          "true tilt in the picture, they just used one fixed, eyeballed 20-degree correction for "
                          "every video - cruder than the crease-detection idea Santhosh proposed earlier, so if we "
                          "finish that properly it would actually be more accurate than what's published here. They "
                          "also never attempted to work out the ball's HEIGHT the way we did (bounce position "
                          "only) - our depth-aware height work goes further than this paper did."),
]
for row in new_notes:
    ws2.append(row)
for row_i in range(ws2.max_row - len(new_notes) + 1, ws2.max_row + 1):
    ws2.cell(row=row_i, column=1).font = BASE_FONT
    ws2.cell(row=row_i, column=2).font = BASE_FONT
    ws2.cell(row=row_i, column=2).alignment = Alignment(wrap_text=True)
wb.save('cricket_tracking.xlsx')
print('done')
